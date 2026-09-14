#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""极简 PPTX 生成器（simple-pptx-generator 技能）。

设计取向是「轻量化汇报」：16:9、纯白底、一个强调色、大标题、大留白、
每页最多 5 条。不追求花哨版式，追求能在 30 秒内读完一页。

两种用法
--------
命令行（推荐）：

    python3 /skills/simple-pptx-generator/generate_pptx.py \\
      --topic "2026 Q1 业务复盘" \\
      --points "市场概况,增长驱动：用户增长；渠道扩张,风险：供应链；合规" \\
      --page-num 8 --data "12,25,31,44" \\
      --output /workspace/q1_review.pptx --json

当函数调用（方式二，可自行改版式）：

    from generate_pptx import generate_light_pptx
    generate_light_pptx(topic="...", points=[...], page_num=8, output="/workspace/x.pptx")

依赖
----
python-pptx。沙箱里通常已由 MINI-RAG 的 provisioning 预装；万一没有，本脚本会自己
补装（含 PEP 668 所需的 --break-system-packages），所以直接跑就行。
"""

from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import time
from typing import Iterable, List, Optional, Sequence

# --------------------------------------------------------------------------- #
# 依赖：缺了自己装，避免"技能不可用"要人肉排查
# --------------------------------------------------------------------------- #

_PIP_INDEX = "https://pypi.tuna.tsinghua.edu.cn/simple"


def _log(msg: str) -> None:
    """进度一律走 stderr —— --json 模式下 stdout 必须只有 JSON。"""
    print(f"[pptx] {msg}", file=sys.stderr, flush=True)


def _pip_works(python: str) -> bool:
    try:
        return subprocess.run(
            [python, "-m", "pip", "--version"],
            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=60,
        ).returncode == 0
    except Exception:
        return False


def _bootstrap_pip(python: str) -> bool:
    """用 get-pip.py 装 pip。Ubuntu 24.04 需要 --break-system-packages（PEP 668）。"""
    import urllib.request

    target = "/tmp/get-pip.py"
    try:
        _log("未找到 pip，正在下载 get-pip.py …")
        with urllib.request.urlopen("https://bootstrap.pypa.io/get-pip.py", timeout=60) as resp:
            data = resp.read()
        with open(target, "wb") as fh:
            fh.write(data)
    except Exception as exc:  # noqa: BLE001
        _log(f"下载 get-pip.py 失败：{exc}")
        return False

    for args in ([python, target, "--quiet", "--break-system-packages"],
                 [python, target, "--quiet"]):
        try:
            if subprocess.run(args, timeout=300).returncode == 0 and _pip_works(python):
                return True
        except Exception:  # noqa: BLE001
            continue
    return False


def ensure_python_pptx() -> None:
    """确保 import pptx 可用；不可用就现场补装。"""
    try:
        import pptx  # noqa: F401
        return
    except ImportError:
        pass

    python = sys.executable or "python3"
    _log("未检测到 python-pptx，准备自动安装（首次约 20–60 秒）…")

    if not _pip_works(python) and not _bootstrap_pip(python):
        raise SystemExit(
            "自动安装失败：无法准备 pip。请在沙箱里手动执行：\n"
            "  curl -fsSL https://bootstrap.pypa.io/get-pip.py -o /tmp/get-pip.py\n"
            "  python3 /tmp/get-pip.py --quiet --break-system-packages\n"
            "  python3 -m pip install --break-system-packages python-pptx"
        )

    base = [python, "-m", "pip", "install", "--quiet", "--disable-pip-version-check",
            "--break-system-packages", "python-pptx"]
    for extra in (["-i", _PIP_INDEX], []):
        try:
            if subprocess.run(base + extra, timeout=600).returncode == 0:
                break
        except Exception as exc:  # noqa: BLE001
            _log(f"安装失败（{extra or '默认源'}）：{exc}")
    else:
        raise SystemExit("自动安装 python-pptx 失败，请检查沙箱网络。")

    try:
        import pptx  # noqa: F401
        _log(f"python-pptx {pptx.__version__} 安装完成")
    except ImportError as exc:
        raise SystemExit(f"安装后仍无法 import pptx：{exc}") from exc


# --------------------------------------------------------------------------- #
# 版式常量：改这里就能整体调风格
# --------------------------------------------------------------------------- #

SLIDE_W = 13.333          # 英寸，16:9
SLIDE_H = 7.5
MARGIN = 0.9
CONTENT_W = SLIDE_W - MARGIN * 2

INK = (0x1F, 0x29, 0x33)      # 正文/标题
MUTED = (0x7B, 0x87, 0x94)    # 次级信息
ACCENT = (0xC6, 0x3A, 0x2B)   # 强调色（与前端的朱红一致）
RULE = (0xE4, 0xE7, 0xEB)     # 分隔线
CHART = (0x2F, 0x4A, 0x6D)    # 图表（靛青）

FONT_CN = "微软雅黑"
FONT_EN = "Segoe UI"

TITLE_SZ = 40
HEADING_SZ = 30
BODY_SZ = 18
SMALL_SZ = 12
MAX_BULLETS = 5


# --------------------------------------------------------------------------- #
# 小工具
# --------------------------------------------------------------------------- #

def _rgb(color: Sequence[int]):
    from pptx.dml.color import RGBColor
    return RGBColor(color[0], color[1], color[2])


def _set_font(run, size: int, color: Sequence[int], bold: bool = False) -> None:
    """设置字号/颜色/加粗，并把东亚字体一并写上，否则中文会走主题默认字体。"""
    from pptx.util import Pt
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)
    run.font.name = FONT_EN
    try:  # a:ea 需要直接改 XML；python-pptx 没有公开 API
        rpr = run._r.get_or_add_rPr()
        from pptx.oxml.ns import qn
        ea = rpr.find(qn("a:ea"))
        if ea is None:
            ea = rpr.makeelement(qn("a:ea"), {})
            rpr.append(ea)
        ea.set("typeface", FONT_CN)
    except Exception:  # noqa: BLE001  拿不到就退回主题字体，不该因此失败
        pass


def _textbox(slide, left, top, width, height):
    from pptx.util import Inches
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    return tf


def _write(tf, text: str, size: int, color, bold=False, align=None,
           line_spacing=None, space_after=None):
    """写入一个段落到文本框：空框用第一段，否则追加。"""
    from pptx.util import Pt
    first = tf.paragraphs[0]
    para = first if (not first.runs and not first.text) else tf.add_paragraph()
    run = para.add_run()
    run.text = text
    _set_font(run, size, color, bold)
    if align is not None:
        para.alignment = align
    if line_spacing is not None:
        para.line_spacing = line_spacing
    if space_after is not None:
        para.space_after = Pt(space_after)
    return para


def _rule(slide, top: float, height: float = 0.02, left: float = MARGIN,
          width: float = CONTENT_W, color=RULE):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top),
                                   Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(color)
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])  # 6 = Blank


def _notes(slide, text: str) -> None:
    if text:
        try:
            slide.notes_slide.notes_text_frame.text = text
        except Exception:  # noqa: BLE001
            pass


def _page_number(slide, number: int) -> None:
    from pptx.enum.text import PP_ALIGN
    tf = _textbox(slide, SLIDE_W - MARGIN - 1.5, SLIDE_H - 0.75, 1.5, 0.35)
    _write(tf, f"{number:02d}", SMALL_SZ, MUTED, align=PP_ALIGN.RIGHT)


# --------------------------------------------------------------------------- #
# 各页版式
# --------------------------------------------------------------------------- #

def _cover(prs, topic: str, subtitle: str, footer: str) -> None:
    from pptx.enum.text import PP_ALIGN
    slide = _blank_slide(prs)
    _rule(slide, 2.45, height=0.09, left=MARGIN, width=1.6, color=ACCENT)

    tf = _textbox(slide, MARGIN, 2.75, CONTENT_W, 1.7)
    _write(tf, topic, TITLE_SZ, INK, bold=True, line_spacing=1.15)

    if subtitle:
        tf = _textbox(slide, MARGIN, 4.55, CONTENT_W, 0.7)
        _write(tf, subtitle, BODY_SZ, MUTED, line_spacing=1.4)

    if footer:
        tf = _textbox(slide, MARGIN, 6.45, CONTENT_W, 0.4)
        _write(tf, footer, SMALL_SZ, MUTED)

    _notes(slide, f"开场：一句话说明《{topic}》要解决什么问题、结论是什么。")


def _section(prs, eyebrow: str, heading: str, bullets: Sequence[str],
             page_no: int, note: str = "") -> None:
    """通用内容页：小标签 + 大标题 + 分隔线 + 条目。"""
    slide = _blank_slide(prs)

    if eyebrow:
        tf = _textbox(slide, MARGIN, 0.72, CONTENT_W, 0.35)
        _write(tf, eyebrow, SMALL_SZ, ACCENT, bold=True)

    tf = _textbox(slide, MARGIN, 1.12, CONTENT_W, 1.0)
    _write(tf, heading, HEADING_SZ, INK, bold=True, line_spacing=1.2)

    _rule(slide, 2.22)

    if bullets:
        tf = _textbox(slide, MARGIN, 2.55, CONTENT_W, 3.9)
        first = True
        for item in bullets[:MAX_BULLETS]:
            para = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            run = para.add_run()
            run.text = f"—  {item}"
            _set_font(run, BODY_SZ, INK)
            para.line_spacing = 1.5
            from pptx.util import Pt
            para.space_after = Pt(14)

    _page_number(slide, page_no)
    _notes(slide, note or heading)


def _agenda(prs, items: Sequence[str], page_no: int) -> None:
    from pptx.util import Pt
    slide = _blank_slide(prs)
    tf = _textbox(slide, MARGIN, 0.72, CONTENT_W, 0.35)
    _write(tf, "AGENDA", SMALL_SZ, ACCENT, bold=True)
    tf = _textbox(slide, MARGIN, 1.12, CONTENT_W, 0.9)
    _write(tf, "目录", HEADING_SZ, INK, bold=True)
    _rule(slide, 2.22)

    tf = _textbox(slide, MARGIN, 2.6, CONTENT_W, 3.8)
    for i, item in enumerate(items[:MAX_BULLETS]):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = para.add_run()
        run.text = f"{i + 1:02d}   {item}"
        _set_font(run, BODY_SZ, INK)
        para.line_spacing = 1.5
        para.space_after = Pt(14)

    _page_number(slide, page_no)


def _chart_page(prs, eyebrow: str, heading: str, values: Sequence[float],
                captions: Sequence[str], page_no: int) -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
    from pptx.util import Inches

    slide = _blank_slide(prs)
    tf = _textbox(slide, MARGIN, 0.72, CONTENT_W, 0.35)
    _write(tf, eyebrow, SMALL_SZ, ACCENT, bold=True)
    tf = _textbox(slide, MARGIN, 1.12, CONTENT_W, 0.9)
    _write(tf, heading, HEADING_SZ, INK, bold=True)
    _rule(slide, 2.22)

    data = CategoryChartData()
    if captions and len(captions) == len(values):
        data.categories = [str(c) for c in captions]
    else:
        data.categories = [f"{i + 1}" for i in range(len(values))]
    data.add_series("数值", tuple(float(v) for v in values))

    frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(MARGIN), Inches(2.5), Inches(CONTENT_W), Inches(3.9), data,
    )
    chart = frame.chart
    chart.has_legend = False
    chart.has_title = False
    try:
        plot = chart.plots[0]
        plot.has_data_labels = True
        plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        plot.data_labels.font.size = _pt(12)
        plot.gap_width = 120
        series = plot.series[0]
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = _rgb(CHART)
    except Exception:  # noqa: BLE001  图表细节随版本变动，失败不影响出图
        pass

    _page_number(slide, page_no)
    _notes(slide, f"数据页：先给结论，再解释 {heading} 的数值含义与口径。")


def _pt(size: int):
    from pptx.util import Pt
    return Pt(size)


def _closing(prs, text: str, sub: str, page_no: int) -> None:
    from pptx.enum.text import PP_ALIGN
    slide = _blank_slide(prs)
    tf = _textbox(slide, MARGIN, 3.0, CONTENT_W, 1.1)
    _write(tf, text, TITLE_SZ, INK, bold=True, align=PP_ALIGN.CENTER)
    if sub:
        tf = _textbox(slide, MARGIN, 4.25, CONTENT_W, 0.6)
        _write(tf, sub, BODY_SZ, MUTED, align=PP_ALIGN.CENTER)
    _page_number(slide, page_no)
    _notes(slide, "收尾：重复一遍核心结论，给出下一步动作。")


# --------------------------------------------------------------------------- #
# 输入解析
# --------------------------------------------------------------------------- #

_SPLIT_POINT = re.compile(r"[,\n，]+")
_SPLIT_BULLET = re.compile(r"[;；\n]+")
_TITLE_SEP = re.compile(r"[:：]")


def _split_points(raw: str) -> List[str]:
    """拆论点。

    只用逗号/换行分隔论点——分号留给页内条目（见 :func:`_parse_point`）。
    两边都用分号会让 ``增长驱动：用户增长；渠道扩张`` 被误拆成两个论点。
    """
    return [p.strip() for p in _SPLIT_POINT.split(raw or "") if p.strip()]


def _parse_point(point: str) -> "tuple[str, List[str]]":
    """``标题：要点1；要点2`` → ``("标题", ["要点1", "要点2"])``。

    没有分隔符时整条当标题，作为一页的论点。
    """
    m = _TITLE_SEP.search(point)
    if not m:
        return point.strip(), []
    title = point[:m.start()].strip()
    rest = point[m.end():].strip()
    bullets = [b.strip() for b in _SPLIT_BULLET.split(rest) if b.strip()]
    if not title:                      # 以冒号开头，退回整条当标题
        return point.strip(), []
    return title, bullets


def _parse_numbers(raw: str) -> List[float]:
    out: List[float] = []
    for token in re.split(r"[,\s，]+", (raw or "").strip()):
        if not token:
            continue
        try:
            out.append(float(token))
        except ValueError:
            raise SystemExit(f"--data 里有非数字项：{token!r}")
    return out


def _chunk(items: Sequence, groups: int) -> List[List]:
    """把 items 尽量均匀地分成 groups 组（元素类型不限）。"""
    groups = max(1, min(groups, len(items)))
    size, extra = divmod(len(items), groups)
    out, cursor = [], 0
    for i in range(groups):
        take = size + (1 if i < extra else 0)
        out.append(list(items[cursor:cursor + take]))
        cursor += take
    return [g for g in out if g]


# --------------------------------------------------------------------------- #
# 主入口
# --------------------------------------------------------------------------- #

def generate_light_pptx(
    topic: str,
    points: Optional[Iterable[str]] = None,
    page_num: int = 8,
    output: str = "/workspace/light_presentation.pptx",
    data: Optional[Sequence[float]] = None,
    subtitle: str = "",
    footer: str = "",
    closing: str = "谢谢",
) -> str:
    """生成一份极简 PPTX，返回输出文件路径。

    参数
    ----
    topic:    封面主题（必填）
    points:   论点列表；每条可以是 ``标题``，也可以是 ``标题：要点1；要点2``
    page_num: 目标总页数（上限）。实际页数由内容决定，不会为空凑页数
    output:   输出路径，建议放 ``/workspace/`` 下
    data:     可选，柱状图数值（会额外生成一页数据页）
    subtitle/footer: 封面的副标题与页脚
    closing:  收尾页大字
    """
    ensure_python_pptx()

    from pptx import Presentation
    from pptx.util import Inches

    topic = (topic or "").strip()
    if not topic:
        raise SystemExit("--topic 不能为空")

    parsed = [_parse_point(p) for p in (points or [])]
    parsed = [(t, b) for t, b in parsed if t]
    values = list(data or [])

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    page = 1
    _cover(prs, topic, subtitle or time.strftime("%Y-%m-%d"), footer)
    page += 1

    use_agenda = len(parsed) >= 4 and page_num >= 5
    if use_agenda:
        _agenda(prs, [t for t, _ in parsed], page)
        page += 1

    # 收尾页 + 目录页 + 数据页都占位，剩下的才是内容页配额
    reserved = 1 + (1 if use_agenda else 0) + (1 if values else 0)
    slots = max(1, page_num - reserved - 1)
    if len(parsed) <= slots:
        sections: List["tuple[str, List[str]]"] = [(t, list(b)) for t, b in parsed]
    else:
        # 论点比页数多：合并。每页标题取该组第一个论点的标题，其余论点压成条目——
        # 注意别把第一个论点再当条目写一遍，否则页标题和第一条会重复。
        sections = []
        for group in _chunk(parsed, slots):
            head, first_bullets = group[0]
            bullets: List[str] = list(first_bullets)
            for t, bs in group[1:]:
                bullets.append(f"{t}：{'；'.join(bs)}" if bs else t)
            sections.append((head, bullets[:MAX_BULLETS]))

    for i, (heading, bullets) in enumerate(sections, start=1):
        _section(prs, f"要点 {i:02d}", heading, bullets, page)
        page += 1

    if values:
        _chart_page(prs, "数据", "关键指标", values, [], page)
        page += 1

    _closing(prs, closing, topic, page)

    out_path = output
    parent = os.path.dirname(os.path.abspath(out_path))
    if parent and not os.path.isdir(parent):
        os.makedirs(parent, exist_ok=True)
    prs.save(out_path)
    return out_path


def _slide_summary(path: str) -> List[dict]:
    """读回生成的文件，给出每页的标题，便于 --json 消费。"""
    from pptx import Presentation
    out = []
    try:
        prs = Presentation(path)
        for i, slide in enumerate(prs.slides, start=1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        t = "".join(r.text for r in para.runs).strip()
                        if t:
                            texts.append(t)
                if getattr(shape, "has_chart", False):
                    texts.append("[chart]")
            title = texts[0] if texts else ""
            out.append({"index": i, "title": title, "texts": texts})
    except Exception:  # noqa: BLE001
        pass
    return out


def main(argv: Optional[Sequence[str]] = None) -> int:
    ap = argparse.ArgumentParser(
        prog="generate_pptx.py",
        description="极简 PPTX 生成器（16:9、纯白、单一强调色、每页最多 5 条）",
    )
    ap.add_argument("--topic", required=True, help="封面主题")
    ap.add_argument("--points", default="",
                    help="论点，逗号分隔；每条可写成「标题：要点1；要点2」")
    ap.add_argument("--page-num", type=int, default=8, help="目标总页数（上限，默认 8）")
    ap.add_argument("--output", default="/workspace/light_presentation.pptx", help="输出 pptx 路径")
    ap.add_argument("--data", default="", help="可选，柱状图数值，如 \"12,25,31,44\"")
    ap.add_argument("--subtitle", default="", help="封面副标题（默认当天日期）")
    ap.add_argument("--footer", default="", help="封面页脚")
    ap.add_argument("--closing", default="谢谢", help="收尾页大字")
    ap.add_argument("--json", action="store_true", help="以 JSON 输出结果（stdout 只有 JSON）")
    args = ap.parse_args(argv)

    points = _split_points(args.points)

    def _fail(message: str) -> int:
        """--json 模式下失败也必须输出 JSON，否则调用方只能看到一个空 stdout。"""
        if args.json:
            print(json.dumps({"ok": False, "error": message}, ensure_ascii=False))
        else:
            print(f"失败：{message}", file=sys.stderr)
        return 1

    try:
        data = _parse_numbers(args.data)
        out = generate_light_pptx(
            topic=args.topic,
            points=points,
            page_num=max(2, args.page_num),
            output=args.output,
            data=data,
            subtitle=args.subtitle,
            footer=args.footer,
            closing=args.closing,
        )
    except SystemExit as exc:          # 参数校验 / 依赖安装失败都会走这里
        return _fail(str(exc) or "参数或环境不满足要求")
    except Exception as exc:  # noqa: BLE001
        return _fail(f"{type(exc).__name__}: {exc}")

    slides = _slide_summary(out)
    result = {
        "ok": True,
        "output": out,
        "pages": len(slides) or None,
        "bytes": os.path.getsize(out) if os.path.exists(out) else 0,
        "slides": [{"index": s["index"], "title": s["title"]} for s in slides],
    }
    if args.json:
        print(json.dumps(result, ensure_ascii=False))
    else:
        print(f"已生成：{out}")
        print(f"共 {result['pages']} 页，{result['bytes']} 字节")
        for s in result["slides"]:
            print(f"  {s['index']:>2}. {s['title']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
