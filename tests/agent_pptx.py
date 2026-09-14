"""端到端验收：agent 是否真的会用 simple-pptx-generator 技能做出 PPT。

链路：用户要 PPT → agent 读 SKILL.md → 在沙箱里跑 generate_pptx.py
      → 回复里给出 /api/sandbox/download 链接 → 该链接能下到合法 pptx

运行前请确保 app 在 8088 且沙箱可用。会真实调用 DeepSeek。
"""

from __future__ import annotations

import json
import os
import re
import sys
import time

import httpx

BASE = os.getenv("BASE", "http://127.0.0.1:8088")
ORG = "default-org"
USER = "pptx-e2e"

REQUEST = (
    "帮我把下面这份材料做成一份 PPT，主题「2026 Q1 业务复盘」，控制在 8 页左右：\n"
    "1. 市场概况：整体规模持平，竞争加剧\n"
    "2. 增长驱动：用户增长 18%；渠道扩张到 12 个城市\n"
    "3. 风险：供应链集中度高；合规成本上升\n"
    "4. 下一步：Q2 聚焦高毛利产品线\n"
    "有季度营收数据 12、25、31、44（单位百万），请放到图表页。"
)

PASS, FAIL = [], []


def check(name, cond, detail=""):
    (PASS if cond else FAIL).append(name)
    print(("PASS  " if cond else "FAIL  ") + name + (("   [" + str(detail) + "]") if detail else ""))


def main():
    payload = {"message": REQUEST, "thread_id": f"{ORG}__{USER}__{int(time.time())}",
               "user_id": USER, "org_id": ORG}

    tool_calls, tool_results, reply = [], [], ""
    t0 = time.time()
    with httpx.Client(timeout=httpx.Timeout(600.0, read=600.0), ) as c:
        with c.stream("POST", BASE + "/api/chat/stream", json=payload) as resp:
            resp.raise_for_status()
            buf = ""
            for chunk in resp.iter_text():
                buf += chunk
                while "\n\n" in buf:
                    frame, buf = buf.split("\n\n", 1)
                    for line in frame.split("\n"):
                        if not line.startswith("data:"):
                            continue
                        try:
                            d = json.loads(line[5:].strip())
                        except json.JSONDecodeError:
                            continue
                        t = d.get("type")
                        if t == "tool_call":
                            tool_calls.append(d.get("name"))
                            print(f"  [{time.time()-t0:5.1f}s] tool_call  {d.get('name')}")
                        elif t == "tool_result":
                            tool_results.append(d.get("output") or "")
                        elif t == "done":
                            reply = d.get("reply") or reply
    print(f"\n耗时 {time.time()-t0:.1f}s，工具调用 {len(tool_calls)} 次")
    print("回复:", reply[:400].replace("\n", " "))

    check("agent 调用了 read_file（读技能文档）", "read_file" in tool_calls, tool_calls)
    check("agent 在沙箱里执行了命令", "execute" in tool_calls, tool_calls)
    # execute 的结果是脚本 stdout（也就是那段 JSON），不含命令本身
    check("技能脚本真的跑出了结果",
          any('"ok": true' in r and ".pptx" in r for r in tool_results),
          next((r[:90] for r in tool_results if ".pptx" in r), ""))
    check("回复里给出了下载链接",
          "/api/sandbox/download" in reply and "path=" in reply, reply[-160:].replace("\n", " "))

    # agent 有时写成 markdown 链接 `[...](url)`、有时用反引号包住，格式每轮都可能不同，
    # 所以先粗抓再剥掉包裹字符。
    m = re.search(r"path=([^\s)\]&\"'<>`]+)", reply)
    if not m:
        check("能从回复里解析出 pptx 路径", False, "未找到 path 参数")
        return finish()

    path = m.group(1).strip("`'\"<>")
    check("能从回复里解析出 pptx 路径", path.endswith(".pptx"), path)

    r = httpx.get(BASE + "/api/sandbox/download", params={"org_id": ORG, "path": path}, timeout=120)
    check("下载链接可用（200）", r.status_code == 200, r.status_code)
    if r.status_code != 200:
        return finish()

    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "out", "agent_made.pptx")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    with open(out, "wb") as fh:
        fh.write(r.content)
    check("下载体积合理（>20KB）", len(r.content) > 20_000, f"{len(r.content)} bytes")

    try:
        from pptx import Presentation
        prs = Presentation(out)
        slides = list(prs.slides)
        title = "".join(run.text for sh in slides[0].shapes if sh.has_text_frame
                        for para in sh.text_frame.paragraphs for run in para.runs)
        check("产物是合法 pptx 且页数合理", len(slides) >= 4, f"{len(slides)} 页")
        check("封面是要求的主题", "2026 Q1" in title, title[:60].replace("\n", " "))
        check("含图表页", any(getattr(sh, "has_chart", False) for s in slides for sh in s.shapes))
    except Exception as exc:  # noqa: BLE001
        check("产物能被 python-pptx 打开", False, f"{type(exc).__name__}: {exc}")

    return finish()


def finish():
    print(f"\n{len(PASS)}/{len(PASS) + len(FAIL)} checks passed")
    if FAIL:
        print("failed: " + "; ".join(FAIL))
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
